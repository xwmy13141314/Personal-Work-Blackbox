"""本地文档文字读取：仅在用户选择后读取，供待办智能提取使用。"""

from __future__ import annotations

import csv
from pathlib import Path


SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".docx", ".pdf"}
MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_TEXT_LENGTH = 100_000


class LocalDocumentError(ValueError):
    """用户选择的本地文档无法安全读取时抛出。"""


def extract_local_document_text(path: str | Path) -> str:
    """读取受支持文档的可提取文字，不上传或保留原始文件。"""
    document = Path(path).expanduser().resolve()
    if not document.is_file():
        raise LocalDocumentError("所选文件不存在或不是普通文件")
    if document.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise LocalDocumentError("仅支持 TXT、Markdown、CSV、Word（DOCX）和 PDF 文件")
    if document.stat().st_size > MAX_FILE_SIZE:
        raise LocalDocumentError("单个文件不能超过 5 MB")

    suffix = document.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        text = _read_text(document)
    elif suffix == ".csv":
        text = _read_csv(document)
    elif suffix == ".docx":
        text = _read_docx(document)
    else:
        text = _read_pdf(document)

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise LocalDocumentError("未能从该文档读取文字；扫描版 PDF 请先进行 OCR")
    return text[:MAX_TEXT_LENGTH]


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeError:
            continue
    raise LocalDocumentError("文档编码无法识别，请另存为 UTF-8 后重试")


def _read_csv(path: Path) -> str:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            return "\n".join(" | ".join(cell.strip() for cell in row if cell.strip()) for row in csv.reader(file))
    except UnicodeError:
        with path.open("r", encoding="gb18030", newline="") as file:
            return "\n".join(" | ".join(cell.strip() for cell in row if cell.strip()) for row in csv.reader(file))


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # 打包漏依赖时给用户明确提示
        raise LocalDocumentError("Word 文档读取组件未安装") from exc
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # 打包漏依赖时给用户明确提示
        raise LocalDocumentError("PDF 文档读取组件未安装") from exc
    reader = PdfReader(path)
    return "\n".join((page.extract_text() or "") for page in reader.pages[:100])
